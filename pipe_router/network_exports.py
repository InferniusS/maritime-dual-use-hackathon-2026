"""Revision-bound designer bundles with independently checked coordinates.

These files are implementation instructions and previews. Writing a bundle never
changes a CAD drawing and never grants approval to replace native solids.
"""
from __future__ import annotations

from copy import deepcopy
import csv
from html import escape
import hashlib
import json
import math
from pathlib import Path
import tempfile
from typing import Any

from .network_manufacturing import manufacture_network
from .network_components import component_geometry

_EXPORT_FILES = {
    "instructions": "instructions.json", "instructions_markdown": "instructions.md",
    "xlsx": "components.xlsx", "csv": "components.csv",
    "dxf": "network-preview.dxf", "result": "network-result.json",
    "comparison": "network-comparison.svg", "coordinates_csv": "coordinates.csv",
}
_COMPONENT_COLUMNS = ["Component ID", "Original CAD ID", "Piping system", "Component type",
    "Nominal diameter mm", "Fitting specification", "Quantity", "Global X", "Global Y", "Global Z",
    "Units", "Installation notes", "Verification status"]
_COORDINATE_COLUMNS = ["Route state", "Component ID", "Original CAD ID", "Piping system", "Point index",
                       "Global X", "Global Y", "Global Z", "Units", "Drawing revision"]


def _json(value):
    return json.dumps(value, sort_keys=True, ensure_ascii=False, allow_nan=False, indent=2)


def _display(value):
    return "MISSING — UNVERIFIED" if value is None or value == "" else value


def _safe_csv(value):
    # Spreadsheet programs interpret these leading characters as formulas.
    # JSON instructions retain exact identifiers; CSV protects on-open execution.
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'"+value
    return value


def _coordinates(original, candidate):
    records = []
    for state, network in (("original", original), ("proposed", candidate)):
        revision = network.get("drawing", {}).get("revision")
        for span in network.get("spans", []):
            for index, point in enumerate(span["points"]):
                records.append([state, span["id"], _display(span.get("cad_handle")), span.get("system_id"), index,
                                *point, "mm", revision])
    return records


def _components(network, manufacture):
    rows = []
    physical = {item["id"]: item for item in component_geometry(network)}
    for piece in manufacture["pieces"]:
        point = piece["points"][0] if piece["points"] else [None]*3
        rows.append([piece["id"], _display(piece.get("cad_handle")), piece.get("system_id"), "manufacturing piece",
                     _display(piece.get("nominal_diameter_mm")), "See physical fitting schedule; fabrication method requires confirmation", 1,
                     *point, "mm", "; ".join(piece.get("notes", []))+f" Developed length {piece['length_mm']:.9g} mm; bends {piece['bend_count']}", piece["verification_status"]])
    for fitting in manufacture["fittings"]:
        point = fitting.get("intersection", fitting.get("center", [None]*3))
        spec = f"Elbow {fitting.get('angle_deg', 'MISSING')} degrees; centerline radius {fitting.get('radius_mm', 'MISSING')} mm; source {_display(fitting.get('source'))}"
        rows.append([fitting["id"], _display(fitting.get("cad_handle")), fitting.get("system_id"), "elbow requirement",
                     _display(fitting.get("nominal_diameter_mm")), spec, 1, *point, "mm",
                     "Location is centerline tangent intersection; tangent points and arc center are in instructions.json. Catalog selection is not a purchase approval.", fitting["verification_status"]])
    for node in network.get("nodes", []):
        body = physical.get(node["id"])
        specification = _json({"dimensions": body["dimensions"], "axis": body["axis"],
                               "physical_envelope": {key: body[key] for key in ("representation", "min", "max", "capsules", "source", "conservative")}}) if body else _display((node.get("dimensions") or {}).get("specification"))
        rows.append([node["id"], _display(node.get("cad_handle")), node.get("system_id"), node.get("type"),
                     _display(node.get("nominal_diameter_mm")), specification, 1,
                     *node["position"], "mm", ("Fixed connection: preserve global coordinates" if node.get("fixed") else "Movable only within declared permitted region") + "; conservative envelope is not commercial or fabrication approval", "UNVERIFIED"])
    for body in physical.values():
        if body["type"] != "support" or not body["id"].startswith("support:"):
            continue
        rows.append([body["id"], _display(body.get("cad_handle")), _display(body.get("system_id")), "support",
                     _display(body.get("nominal_diameter_mm")), _json({key: body[key] for key in ("dimensions", "axis", "representation", "min", "max", "capsules", "source")}), 1,
                     *(body.get("position") or [None]*3), "mm", "Position basis: "+body["position_basis"]+"; supplied geometry does not establish support load capacity", "UNVERIFIED"])
    for weld in manufacture.get("welds", []):
        point = weld.get("position") or [None]*3
        span = next((span for span in network["spans"] if span["id"] == weld.get("span_id")), {})
        rows.append([weld["id"], _display(weld.get("cad_handle")), weld.get("system_id"), "weld location",
                     _display(span.get("nominal_diameter_mm")), "Weld type, process and consumable require engineering confirmation", 1,
                     *point, "mm", f"Developed station {weld.get('station_mm')} mm; {weld.get('status', 'unverified')}. {weld.get('verification_scope', '')} No welding approval is inferred.",
                     weld.get("verification_status", "UNVERIFIED")])
    return rows


def _check_summary(validation):
    counts = {status: 0 for status in ("PASS", "FAIL", "UNVERIFIED", "NOT_APPLICABLE")}
    for check in validation["checks"]:
        counts[check["status"]] += 1
    return {"status": "FAIL" if validation["hard_failures"] else "UNVERIFIED" if counts["UNVERIFIED"] else "PASS",
            "hard_failures": validation["hard_failures"], "check_counts": counts}


def _component_verification(identifier, validation):
    # A network-wide unknown can affect a component even when no individual
    # object could be identified. Do not turn an empty ID list into a PASS.
    applicable = [check for check in validation["checks"] if not check.get("component_ids") or identifier in check["component_ids"]]
    status = "FAIL" if any(check["status"] == "FAIL" and check["severity"] == "hard" for check in applicable) else "UNVERIFIED" if any(check["status"] == "UNVERIFIED" for check in applicable) else "PASS"
    return {"status": status, "check_ids": [check["id"] for check in applicable],
            "remaining_check_ids": [check["id"] for check in applicable if check["status"] in ("FAIL", "UNVERIFIED")]}


def _dimensions(span):
    from .network_fittings import corner_angle
    points = span["points"]
    return {"length_basis":"Routing polyline through tangent intersections; physical developed bend/piece lengths are in the manufacturing schedule",
            "centerline_length_mm": sum(math.dist(a, b) for a, b in zip(points, points[1:])),
            "direction_changes": sum(corner_angle(a, b, c) > 1e-5 for a, b, c in zip(points, points[1:], points[2:])),
            "segments": [{"index": index, "start": list(a), "end": list(b), "length_mm": math.dist(a, b)}
                         for index, (a, b) in enumerate(zip(points, points[1:]))]}


def _segment_difference(points, retained_points):
    retained = {(tuple(a), tuple(b)) for a, b in zip(retained_points, retained_points[1:])}
    return [{"index": index, "start": list(a), "end": list(b), "length_mm": math.dist(a, b)}
            for index, (a, b) in enumerate(zip(points, points[1:])) if (tuple(a), tuple(b)) not in retained]


def _network_audit(source, proposed, original_validation, validation):
    originals = {span["id"]: span for span in source["spans"]}
    changes = []
    for span in proposed["spans"]:
        old = originals[span["id"]]
        before, after = _dimensions(old), _dimensions(span)
        changed = old["points"] != span["points"]
        length_delta = after["centerline_length_mm"]-before["centerline_length_mm"]
        bend_delta = after["direction_changes"]-before["direction_changes"]
        changes.append({"span_id": span["id"], "cad_handle": span.get("cad_handle"), "system_id": span["system_id"],
                        "state": "changed" if changed else "untouched", "nominal_diameter_mm": span.get("nominal_diameter_mm"),
                        "outer_diameter_mm": span.get("outer_diameter_mm"), "from_node": span["from_node"], "to_node": span["to_node"],
                        "original_start": list(old["points"][0]), "original_end": list(old["points"][-1]),
                        "proposed_start": list(span["points"][0]), "proposed_end": list(span["points"][-1]),
                        "before": before, "after": after, "length_delta_mm": length_delta, "bend_delta": bend_delta,
                        "before_validation": _component_verification(span["id"], original_validation),
                        "after_validation": _component_verification(span["id"], validation),
                        "rationale": f"Calculated centerline length {before['centerline_length_mm']:.6f} to {after['centerline_length_mm']:.6f} mm ({length_delta:+.6f} mm); direction changes {before['direction_changes']} to {after['direction_changes']}. "
                                     + ("Review this change together with every other changed span; separate per-span implementation is not approved." if changed else "Existing route geometry is preserved and checked in the final network environment.")})
    systems = []
    for system_id in sorted({span["system_id"] for span in proposed["spans"]}):
        rows = [change for change in changes if change["system_id"] == system_id]
        systems.append({"system_id": system_id, "span_ids": [row["span_id"] for row in rows],
                        "changed_span_ids": [row["span_id"] for row in rows if row["state"] == "changed"],
                        "untouched_span_ids": [row["span_id"] for row in rows if row["state"] == "untouched"],
                        "before_length_mm": sum(row["before"]["centerline_length_mm"] for row in rows),
                        "after_length_mm": sum(row["after"]["centerline_length_mm"] for row in rows),
                        "before_direction_changes": sum(row["before"]["direction_changes"] for row in rows),
                        "after_direction_changes": sum(row["after"]["direction_changes"] for row in rows)})
    original_nodes = {node["id"]: node for node in source["nodes"]}
    moved_nodes = [{"node_id": node["id"], "cad_handle": node.get("cad_handle"), "system_id": node.get("system_id"),
                    "original_position": original_nodes[node["id"]]["position"], "proposed_position": node["position"],
                    "displacement_mm": math.dist(original_nodes[node["id"]]["position"], node["position"]),
                    "incident_span_ids": [span["id"] for span in proposed["spans"] if node["id"] in (span["from_node"], span["to_node"])]}
                   for node in proposed["nodes"] if node["position"] != original_nodes[node["id"]]["position"]]
    return {"spans": changes, "systems": systems, "moved_nodes": moved_nodes,
            "changed_span_ids": [row["span_id"] for row in changes if row["state"] == "changed"],
            "untouched_span_ids": [row["span_id"] for row in changes if row["state"] == "untouched"],
            "original_validation": _check_summary(original_validation), "proposed_validation": _check_summary(validation),
            "coordinate_basis": "Exact source global WCS millimetres; identity transform. Display dimensions are rounded labels only."}


def _candidate_comparison(result, source_digest, candidate_digest, before, after, original_validation, validation):
    rows = [{"record_type": "independently_recomputed", "candidate_id": "original", "candidate_digest": source_digest,
             "selected": candidate_digest == source_digest, "metrics": before, "validation": _check_summary(original_validation)},
            {"record_type": "independently_recomputed", "candidate_id": "delivered", "candidate_digest": candidate_digest,
             "selected": True, "metrics": after, "validation": _check_summary(validation)}]
    for attempt in result.get("search", {}).get("attempts", []):
        if "evaluation" not in attempt:
            continue
        rows.append({"record_type": "optimizer_recorded_attempt", "candidate_id": str(attempt["evaluation"]),
                     "candidate_digest": attempt.get("candidate_digest"), "selected": attempt.get("selected", False),
                     "accepted_during_search": attempt.get("accepted", False), "metrics": deepcopy(attempt.get("metrics")),
                     "recorded_objective": attempt.get("objective"), "recorded_hard_failures": attempt.get("hard_failures"),
                     "changed_span_ids": attempt.get("changed_span_ids", []), "routing_order": attempt.get("order", []),
                     "reservation_mode": attempt.get("reservation_mode"),
                     "evidence_scope": "Optimizer search log. Only the original and delivered geometries are independently recalculated during export."})
    return rows


def _installation_steps(source, audit, manufacture, has_proposal):
    if not has_proposal or not audit["changed_span_ids"]:
        return [{"step": 1, "action": "Preserve source drawing and existing geometry; this bundle does not request a modification.",
                 "gate": "diagnostic_or_unchanged", "component_ids": []}]
    return [
        {"step": 1, "action": f"Verify source drawing {source['drawing']['id']} revision {source['drawing']['revision']}, source digest and all CAD handles against the active drawing.", "gate": "required_before_preview_or_approval", "component_ids": audit["changed_span_ids"]},
        {"step": 2, "action": "Review all changed routes as one coordinated network proposal; inspect untouched routes and fixed connections in the final shared environment.", "gate": "designer_review_required", "component_ids": audit["changed_span_ids"]},
        {"step": 3, "action": "Resolve every applicable FAIL and UNVERIFIED item, including missing catalog availability, material, weld process, supports, access and installation sequence. Collision clearance alone does not validate installation feasibility.", "gate": "engineering_resolution_required", "component_ids": audit["changed_span_ids"]},
        {"step": 4, "action": "Confirm physical fitting dimensions, developed piece lengths and weld stations against an approved fabrication plan. Fitting requirements are constituents of piece takeoff and must not be added again as pipe length.", "gate": "approved_fabrication_plan_required", "component_ids": [piece["id"] for piece in manufacture["pieces"] if piece["span_id"] in audit["changed_span_ids"]]},
        {"step": 5, "action": "Preview proposed centerlines, bend envelopes and exact global coordinates on separate review layers. Obtain explicit designer approval before changing any native object.", "gate": "native_apply_not_authorized_by_export", "component_ids": audit["changed_span_ids"]},
        {"step": 6, "action": "After authorized implementation, re-extract the drawing and independently validate the complete resulting network before fabrication or installation release.", "gate": "post_implementation_validation_required", "component_ids": audit["changed_span_ids"]},
    ]


def _write_xlsx(path, components, coordinates, manufacture, validation, metadata, instructions):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    workbook = Workbook()
    workbook.remove(workbook.active)
    def sheet(name, headings, rows):
        worksheet = workbook.create_sheet(name)
        worksheet.append(headings)
        for row in rows:
            # XLSX has a 32767-character cell limit. Never silently lose a long
            # validation log; retain full evidence in the JSON instruction file.
            worksheet.append([value[:32680]+"… [TRUNCATED: full evidence in instructions.json]" if isinstance(value, str) and len(value) > 32767 else value for value in row])
            # Preserve exact CAD/engineering IDs as text, including '=' prefixes.
            for cell in worksheet[worksheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="193D54")
        for index, heading in enumerate(headings, 1):
            worksheet.column_dimensions[get_column_letter(index)].width = min(52, max(16, len(heading)+3))
        return worksheet
    sheet("Components", _COMPONENT_COLUMNS, components)
    sheet("Coordinates", _COORDINATE_COLUMNS, coordinates)
    sheet("Manufacturing", ["Piece ID", "Span ID", "Status", "Developed length mm", "Bend count", "Start station mm", "End station mm", "Verification", "Notes"],
          [[piece["id"], piece["span_id"], piece["status"], piece["length_mm"], piece["bend_count"], piece.get("start_station_mm"), piece.get("end_station_mm"), piece["verification_status"], "; ".join(piece.get("notes", []))] for piece in manufacture["pieces"]])
    sheet("Checks", ["Check ID", "Status", "Severity", "Components", "Message", "Numerical evidence"],
          [[check["id"], check["status"], check["severity"], ", ".join(check.get("component_ids", [])), check["message"], _json(check.get("evidence", {}))] for check in validation["checks"]])
    sheet("Baseline checks", ["Check ID", "Status", "Severity", "Components", "Message", "Numerical evidence"],
          [[check["id"], check["status"], check["severity"], ", ".join(check.get("component_ids", [])), check["message"], _json(check.get("evidence", {}))] for check in instructions["original_validation"]["checks"]])
    sheet("Changes", ["Span ID", "CAD handle", "System", "State", "Before length mm", "After length mm", "Length delta mm", "Before bends", "After bends", "Before status", "After status", "Original start XYZ mm", "Proposed start XYZ mm", "Original end XYZ mm", "Proposed end XYZ mm", "Rationale"],
          [[row["span_id"], _display(row["cad_handle"]), row["system_id"], row["state"], row["before"]["centerline_length_mm"], row["after"]["centerline_length_mm"], row["length_delta_mm"], row["before"]["direction_changes"], row["after"]["direction_changes"], row["before_validation"]["status"], row["after_validation"]["status"], _json(row["original_start"]), _json(row["proposed_start"]), _json(row["original_end"]), _json(row["proposed_end"]), row["rationale"]] for row in instructions["network_audit"]["spans"]])
    sheet("Systems", ["System ID", "Changed spans", "Untouched spans", "Before length mm", "After length mm", "Before bends", "After bends"],
          [[row["system_id"], _json(row["changed_span_ids"]), _json(row["untouched_span_ids"]), row["before_length_mm"], row["after_length_mm"], row["before_direction_changes"], row["after_direction_changes"]] for row in instructions["network_audit"]["systems"]])
    sheet("Welds", ["Weld ID", "Span ID", "CAD handle", "System", "Station mm", "Global X", "Global Y", "Global Z", "Units", "Status", "Verification", "Verification scope"],
          [[weld["id"], weld.get("span_id"), _display(weld.get("cad_handle")), weld.get("system_id"), weld.get("station_mm"), *(weld.get("position") or [None]*3), "mm", weld.get("status", "unverified"), weld.get("verification_status", "UNVERIFIED"), weld.get("verification_scope", "UNVERIFIED")] for weld in manufacture.get("welds", [])])
    sheet("Candidates", ["Candidate ID", "Evidence", "Digest", "Selected", "Objective", "Hard failures", "Length mm", "Bends", "Hydraulic evidence", "Normalized terms", "Routing order"],
          [[row["candidate_id"], row["record_type"], row.get("candidate_digest"), row["selected"], (row.get("metrics") or {}).get("objective", row.get("recorded_objective")), row.get("validation", {}).get("hard_failures", row.get("recorded_hard_failures")), (row.get("metrics") or {}).get("length_mm"), (row.get("metrics") or {}).get("direction_changes"), _json((row.get("metrics") or {}).get("hydraulic_evidence", {})), _json((row.get("metrics") or {}).get("objective_terms", {})), _json(row.get("routing_order", []))] for row in instructions["candidate_comparison"]])
    sheet("Installation", ["Step", "Action", "Required gate", "Components"],
          [[row["step"], row["action"], row["gate"], _json(row["component_ids"])] for row in instructions["installation_steps"]])
    sheet("Metadata", ["Field", "Value"], [[key, _json(value) if isinstance(value, (dict, list)) else value] for key, value in metadata.items()])
    workbook.save(path)


def _write_comparison_svg(path, original, proposed, audit):
    """Portable dimensioned view: no scripts, external URLs or CAD transforms."""
    def project(point):
        return ((point[0]-point[1])*math.sqrt(3)/2, (point[0]+point[1])*.5-point[2])
    all_points = [project(point) for network in (original, proposed) for span in network["spans"] for point in span["points"]]
    low = [min(point[i] for point in all_points) for i in (0, 1)]
    high = [max(point[i] for point in all_points) for i in (0, 1)]
    scale = min(490/max(1, high[0]-low[0]), 380/max(1, high[1]-low[1]))
    changed = set(audit["changed_span_ids"])
    rows = ['<svg xmlns="http://www.w3.org/2000/svg" width="1200" height="620" viewBox="0 0 1200 620" role="img" aria-labelledby="title desc">',
            '<title id="title">PipeForge dimensioned original and proposed network comparison</title>',
            '<desc id="desc">Static isometric preview. Labels are measured 3D centerline segment lengths in millimetres. Exact global coordinates are in the coordinate schedule. This is not a fabrication drawing or approval.</desc>',
            '<rect width="1200" height="620" fill="#f4f7fa"/>',
            '<style>text{font-family:Arial,sans-serif;fill:#153148} .dimension{font-size:10px;paint-order:stroke;stroke:#f4f7fa;stroke-width:3px;stroke-linejoin:round} .route{fill:none;stroke-linejoin:round;stroke-linecap:round}</style>',
            '<text x="30" y="32" font-size="22">PipeForge Edge — coordinated network comparison</text>',
            '<text x="30" y="55" font-size="12">Dimension labels: actual 3D centerline length, mm. Isometric preview; exact WCS coordinates are unchanged.</text>']
    labels = 0
    for panel, (state, network) in enumerate((("ORIGINAL", original), ("PROPOSED", proposed))):
        offset = 30+600*panel
        rows.append(f'<text x="{offset}" y="92" font-size="17">{state}</text>')
        def screen(point):
            x, y = project(point)
            return offset+30+(x-low[0])*scale, 125+(y-low[1])*scale
        for span in network["spans"]:
            color = "#"+hashlib.sha256(span["system_id"].encode("utf-8")).hexdigest()[:6]
            points = [screen(point) for point in span["points"]]
            values = " ".join(f"{x:.6f},{y:.6f}" for x, y in points)
            kind = "changed" if span["id"] in changed else "untouched"
            title = escape(f"{span['id']} / {span['system_id']} / {kind}")
            rows.append(f'<polyline class="route" points="{values}" stroke="{color}" stroke-width="{4 if kind == "changed" else 2}" opacity="{1 if kind == "changed" else .65}"><title>{title}</title></polyline>')
            for (a, b), (left, right) in zip(zip(span["points"], span["points"][1:]), zip(points, points[1:])):
                if labels >= 256:
                    break
                label = escape(f"{span['id']}: {math.dist(a, b):.3f} mm")
                rows.append(f'<text class="dimension" x="{(left[0]+right[0])/2:.6f}" y="{(left[1]+right[1])/2-5:.6f}">{label}</text>')
                labels += 1
        total = sum(math.dist(a, b) for span in network["spans"] for a, b in zip(span["points"], span["points"][1:]))
        rows.append(f'<text x="{offset}" y="540" font-size="13">Centerline length: {total:.3f} mm; spans: {len(network["spans"])}</text>')
    rows.extend([f'<text x="30" y="570" font-size="12">Changed spans: {len(changed)}; untouched spans: {len(audit["untouched_span_ids"])}. Color identifies system; heavier lines identify changed spans.</text>',
                 '<text x="30" y="590" font-size="12">Preview excludes context solids and fitting envelopes. Review full DXF and validation report. Dimension labels limited to 256 segments.</text>', '</svg>'])
    path.write_text("\n".join(rows)+"\n", encoding="utf-8")


def _write_dxf(path, original, proposed, manufacture):
    import ezdxf
    doc = ezdxf.new("R2013")
    doc.units = 4  # INSUNITS millimetres; WCS coordinates are never translated.
    doc.appids.new("PIPEFORGE4")
    for name, color in (("PF4_ORIGINAL", 30), ("PF4_PROPOSED", 3), ("PF4_FITTINGS", 4), ("PF4_CONTEXT", 8), ("PF4_PORTS", 2), ("PF4_DIMENSIONS", 7), ("PF51_COMPONENTS", 6)):
        doc.layers.new(name, dxfattribs={"color": color})
    model = doc.modelspace()
    for state, network in (("original", original), ("proposed", proposed)):
        for span in network.get("spans", []):
            entity = model.add_polyline3d(span["points"], dxfattribs={"layer": "PF4_"+state.upper()})
            entity.set_xdata("PIPEFORGE4", [(1000, state)]+[(1000, span["id"][index:index+255]) for index in range(0, len(span["id"]), 255)])
            for index, (a,b) in enumerate(zip(span["points"], span["points"][1:])):
                mid = tuple((x+y)/2 for x,y in zip(a,b))
                model.add_text(f"{state} {span['id']}:{index} {math.dist(a,b):.3f} mm", dxfattribs={"height": 25, "insert": mid, "layer":"PF4_DIMENSIONS"})
    for fitting in manufacture["fittings"]:
        capsules = fitting.get("capsules", [])
        if capsules:
            points = [capsules[0]["start"]]+[capsule["end"] for capsule in capsules]
            model.add_polyline3d(points, dxfattribs={"layer":"PF4_FITTINGS"})
    for node in proposed.get("nodes", []):
        model.add_point(node["position"], dxfattribs={"layer":"PF4_PORTS"})
        model.add_text(node["id"], dxfattribs={"height": 25, "insert": node["position"], "layer":"PF4_PORTS"})
    for obstacle in proposed.get("obstacles", []):
        low, high = obstacle["min"], obstacle["max"]
        corners = [[high[0] if mask&1 else low[0], high[1] if mask&2 else low[1], high[2] if mask&4 else low[2]] for mask in range(8)]
        for index in range(8):
            for axis in (1,2,4):
                if not index&axis:
                    model.add_line(corners[index], corners[index|axis], dxfattribs={"layer":"PF4_CONTEXT"})
    for body in component_geometry(proposed):
        if body["min"] is None:
            continue
        # Wire bounds deliberately remain review envelopes, not fabricated CAD solids.
        low, high = body["min"], body["max"]
        corners = [[high[0] if mask&1 else low[0], high[1] if mask&2 else low[1], high[2] if mask&4 else low[2]] for mask in range(8)]
        for index in range(8):
            for axis in (1, 2, 4):
                if not index&axis:
                    entity = model.add_line(corners[index], corners[index|axis], dxfattribs={"layer":"PF51_COMPONENTS"})
                    entity.set_xdata("PIPEFORGE4", [(1000, "conservative_component_envelope")]+[(1000, body["id"][part:part+255]) for part in range(0, len(body["id"]), 255)])
    doc.saveas(path)


def _verify_serialized_coordinates(directory, original, proposed, coordinates, components):
    """Read all delivered coordinate formats back from disk independently."""
    import ezdxf
    from openpyxl import load_workbook
    document = json.loads((directory/_EXPORT_FILES["instructions"]).read_text(encoding="utf-8"))
    original_by_id = {span["id"]:span for span in original["spans"]}
    proposed_by_id = {span["id"]:span for span in proposed["spans"]}
    for change in document["spans"]:
        if change["original_centerline_mm"] != original_by_id[change["id"]]["points"] or change["proposed_centerline_mm"] != proposed_by_id[change["id"]]["points"]:
            raise ValueError("Export JSON coordinate consistency failed")
    workbook = load_workbook(directory/_EXPORT_FILES["xlsx"], read_only=True, data_only=False)
    rows = list(workbook["Coordinates"].iter_rows(min_row=2, values_only=True))
    xlsx_components = list(workbook["Components"].iter_rows(min_row=2, values_only=True))
    workbook.close()
    if len(rows) != len(coordinates):
        raise ValueError("Export XLSX coordinate count failed")
    max_error = 0.0
    for actual, expected in zip(rows, coordinates):
        if actual[:5] != tuple(expected[:5]):
            raise ValueError("Export XLSX CAD identity consistency failed")
        for index in (5,6,7):
            max_error = max(max_error, abs(actual[index]-expected[index]))
    with (directory/_EXPORT_FILES["coordinates_csv"]).open("r", encoding="utf-8-sig", newline="") as stream:
        csv_coordinates = list(csv.reader(stream))[1:]
    if len(csv_coordinates) != len(coordinates):
        raise ValueError("Export CSV coordinate count failed")
    for actual, expected in zip(csv_coordinates, coordinates):
        if actual[:5] != [str(_safe_csv(value)) for value in expected[:5]]:
            raise ValueError("Export CSV coordinate identity failed")
        for index in (5,6,7):
            value = float(actual[index])
            if not math.isfinite(value):
                raise ValueError("Export CSV coordinate is nonfinite")
            max_error = max(max_error, abs(value-expected[index]))
    with (directory/_EXPORT_FILES["csv"]).open("r", encoding="utf-8-sig", newline="") as stream:
        component_rows = list(csv.reader(stream))[1:]
    if len(component_rows) != len(components) or len(xlsx_components) != len(components):
        raise ValueError("Export component schedule row count failed")
    for csv_row,xlsx_row,expected in zip(component_rows,xlsx_components,components):
        if csv_row[:4] != [str(_safe_csv(value)) for value in expected[:4]] or xlsx_row[:4] != tuple(expected[:4]):
            raise ValueError("Export component schedule identity failed")
        for index in (7,8,9):
            if expected[index] is None:
                if csv_row[index] or xlsx_row[index] is not None:
                    raise ValueError("Export invented an unknown component coordinate")
            else:
                values = (float(csv_row[index]),xlsx_row[index])
                if not all(math.isfinite(value) for value in values):
                    raise ValueError("Export has nonfinite component coordinates")
                max_error = max(max_error,*(abs(value-expected[index]) for value in values))
    dxf = ezdxf.readfile(directory/_EXPORT_FILES["dxf"])
    polylines = [entity for entity in dxf.modelspace().query("POLYLINE") if entity.dxf.layer in ("PF4_ORIGINAL", "PF4_PROPOSED")]
    expected_lines = [(state, span) for state, network in (("original",original),("proposed",proposed)) for span in network["spans"]]
    if len(polylines) != len(expected_lines) or dxf.units != 4:
        raise ValueError("Export DXF route count or units failed")
    for entity, (state, span) in zip(polylines, expected_lines):
        values = list(entity.points())
        if len(values) != len(span["points"]):
            raise ValueError("Export DXF point count failed")
        xdata = entity.get_xdata("PIPEFORGE4")
        if xdata[0].value != state or "".join(tag.value for tag in xdata[1:]) != span["id"]:
            raise ValueError("Export DXF CAD identity failed")
        for actual, expected in zip(values, span["points"]):
            max_error = max(max_error, math.dist(tuple(actual), expected))
    if max_error > 1e-6:
        raise ValueError("Export coordinate error exceeds 0.000001 mm")
    return {"id":"export.serialized_coordinate_roundtrip", "status":"PASS", "severity":"hard",
            "message":"JSON, XLSX, coordinate CSV and DXF were read back and compared with exact global-WCS routes",
            "evidence":{"max_error_mm":max_error,"tolerance_mm":1e-6,"coordinate_rows":len(rows),"csv_coordinate_rows":len(csv_coordinates),"dxf_units":"mm","coordinate_frame":"WCS"},
            "component_ids":list(proposed_by_id)}


def _markdown(instructions):
    lines = ["# PipeForge connected-routing implementation instructions", "", f"Status: {instructions['status']}",
             f"Source drawing: {instructions['drawing'].get('id')} / revision {instructions['source_revision']}",
             "", "Coordinates are global WCS, in millimetres. Do not move fixed ports. A bundle is not approval to replace native CAD solids.",
             "Verify the drawing revision and object handles before implementation. Review every FAIL and UNVERIFIED item before fabrication.",
             "", "Original routes: PF4_ORIGINAL; proposed routes: PF4_PROPOSED; physical bend centerlines: PF4_FITTINGS; conservative component wire bounds: PF51_COMPONENTS. Dimensions are labelled in mm.", ""]
    audit = instructions["network_audit"]
    lines.extend(["## Whole-network calculation", instructions["engineering_explanation"], "",
                  f"Changed spans: {', '.join(audit['changed_span_ids']) or 'none'}.",
                  f"Untouched spans: {', '.join(audit['untouched_span_ids']) or 'none'}.",
                  f"Original hard failures: {audit['original_validation']['hard_failures']}; proposed hard failures: {audit['proposed_validation']['hard_failures']}.",
                  "", "| System | Changed spans | Untouched spans | Before mm | After mm | Before bends | After bends |",
                  "|---|---:|---:|---:|---:|---:|---:|"])
    for system in audit["systems"]:
        label = system["system_id"].replace("|", "\\|").replace("\n", " ")
        lines.append(f"| {label} | {len(system['changed_span_ids'])} | {len(system['untouched_span_ids'])} | {system['before_length_mm']:.6f} | {system['after_length_mm']:.6f} | {system['before_direction_changes']} | {system['after_direction_changes']} |")
    lines.extend(["", "![Dimensioned original and proposed network preview](network-comparison.svg)",
                  "", "This is a centerline comparison, not a fabrication drawing. Exact WCS coordinates are in Coordinates / coordinates.csv. Fitting envelopes and context are in DXF and the validation evidence.", ""])
    for span in instructions["spans"]:
        change = next(row for row in audit["spans"] if row["span_id"] == span["id"])
        lines.extend([f"## {span['id']} (CAD {_display(span.get('cad_handle'))})", f"System: {span['system_id']}; nominal diameter: {_display(span.get('nominal_diameter_mm'))} mm.",
                      f"Change: {'replace centerline according to coordinates after designer approval' if span['changed'] else 'preserve existing route'}.", change["rationale"],
                      f"Final validation status: {change['after_validation']['status']}.", "", "Original XYZ mm:", "```json", _json(span["original_centerline_mm"]), "```", "Proposed XYZ mm:", "```json", _json(span["proposed_centerline_mm"]), "```", ""])
    lines.extend(["## Coordinated implementation sequence", *[f"{step['step']}. {step['action']} Required gate: `{step['gate']}`." for step in instructions["installation_steps"]],
                  "", "## Evidence and candidate tradeoffs", "Original and delivered metrics/checks are independently recalculated during export. Intermediate candidates are identified as optimizer-recorded search evidence; their full geometry is not revalidated by the exporter.",
                  "No global-optimality proof, commercial approval or installation release is implied. Missing hydraulic inputs remain UNVERIFIED; no pressure-loss reduction is invented.",
                  "", "```json", _json(instructions["candidate_comparison"]), "```",
                  "", "## Remaining engineering decisions", *[f"- {_json(issue) if isinstance(issue, (dict, list)) else issue}" for issue in instructions["unresolved_issues"]], "", "See components.xlsx for component, coordinate, manufacturing, weld, change, system, candidate, installation, baseline/final-check and metadata schedules."])
    return "\n".join(lines)+"\n"


def export_network(network: dict[str, Any], result: dict[str, Any], output_dir) -> dict[str, str]:
    """Write a fresh, independently validated instruction bundle; never apply CAD.

    The source digest and revision are mandatory. Supplied validation and metric
    claims are not accepted as evidence. A hard-invalid candidate is rejected.
    No-improvement/blocked reports can be exported, labelled for diagnosis only.
    """
    from .connected import network_digest, metrics
    from .network_model import normalize_network
    from .network_validation import validate_network, validation_delta
    source = normalize_network(network)
    source_digest = network_digest(source)
    revision = source.get("drawing", {}).get("revision")
    if not revision or result.get("source_revision") != revision:
        raise ValueError("Stale result: source drawing revision changed or is missing")
    if result.get("source_digest") != source_digest:
        raise ValueError("Stale result: source network digest changed")
    original_in_result = result.get("original_network")
    if original_in_result is not None and network_digest(normalize_network(original_in_result)) != source_digest:
        raise ValueError("Result original network does not match the current drawing snapshot")
    candidate_raw = result.get("proposed_network")
    candidate = normalize_network(candidate_raw) if candidate_raw is not None else deepcopy(source)
    if candidate.get("drawing") != source.get("drawing") or candidate.get("coordinate_system") != source.get("coordinate_system"):
        raise ValueError("Candidate drawing identity or coordinate frame changed")
    if {span["id"] for span in candidate["spans"]} != {span["id"] for span in source["spans"]}:
        raise ValueError("Candidate component identities changed")
    selection = result.get("selection")
    permitted_span_ids = None
    if selection is not None:
        if not isinstance(selection, dict) or "modifiable_span_ids" not in selection or selection["modifiable_span_ids"] is None:
            raise ValueError("Result selection requires an explicit modifiable_span_ids scope")
        permitted_span_ids = selection["modifiable_span_ids"]
    original_validation = validate_network(source, reference=source, permitted_span_ids=permitted_span_ids)
    validation = validate_network(candidate, reference=source, permitted_span_ids=permitted_span_ids)
    if candidate_raw is not None and validation["hard_failures"]:
        raise ValueError("Candidate failed independent engineering validation; export denied")
    manufacture = manufacture_network(candidate)
    objective_config = result.get("objective_config") or {}
    if not isinstance(objective_config, dict):
        raise ValueError("Objective configuration must be an object")
    metric_options = {"strategy": objective_config.get("strategy", result.get("strategy")),
                      "objective_weights": objective_config.get("weights")}
    before = metrics(source, reference=source, validation=original_validation, **metric_options)
    after = metrics(candidate, reference=source, validation=validation, **metric_options)
    audit = _network_audit(source, candidate, original_validation, validation)
    candidate_digest = network_digest(candidate)
    comparison = _candidate_comparison(result, source_digest, candidate_digest, before, after, original_validation, validation)
    explanation = (f"Calculated complete-network centerline length {before['length_mm']:.6f} to {after['length_mm']:.6f} mm "
                   f"({after['length_mm']-before['length_mm']:+.6f} mm), and direction changes {before['direction_changes']} to {after['direction_changes']}. "
                   f"Changed spans: {len(audit['changed_span_ids'])}; untouched spans rechecked in the final environment: {len(audit['untouched_span_ids'])}. "
                   f"Independent hard failures: {original_validation['hard_failures']} to {validation['hard_failures']}. "
                   "Original and delivered scores use one declared objective configuration and the immutable source network for normalization. "
                   "This comparison provides no proof of global optimality. Unverified checks remain unresolved engineering decisions.")
    original_by_id = {span["id"]:span for span in source["spans"]}
    nodes_by_id = {node["id"]:node for node in candidate["nodes"]}
    physical = component_geometry(candidate)
    unresolved = list(result.get("unresolved_issues", []))
    for check in validation["checks"]:
        if check["status"] in ("FAIL", "UNVERIFIED"):
            message = f"{check['status']}: {check['id']}: {check['message']}"
            if message not in unresolved:
                unresolved.append(message)
    spans = []
    for span in candidate["spans"]:
        original = original_by_id[span["id"]]
        spans.append({"id":span["id"], "pipe_id":span.get("pipe_id", span["id"]), "cad_handle":span.get("cad_handle"), "instance_path":span.get("instance_path", []),
                      "source_instances":deepcopy(span.get("source_instances", [])), "source_layers":deepcopy(span.get("source_layers", [])),
                      "system_id":span["system_id"], "changed":span["points"] != original["points"],
                      "original_centerline_mm":original["points"], "proposed_centerline_mm":span["points"],
                      "from_node":span["from_node"], "to_node":span["to_node"],
                      "fixed_connections":[{"id":identifier,"position":nodes_by_id[identifier]["position"],"cad_handle":nodes_by_id[identifier].get("cad_handle")} for identifier in (span["from_node"],span["to_node"]) if nodes_by_id[identifier].get("fixed")],
                      "nominal_diameter_mm":span.get("nominal_diameter_mm"), "outer_diameter_mm":span.get("outer_diameter_mm"),
                      "dimensions":span.get("dimensions", {}), "constraints":span.get("constraints", {}),
                      "original_dimensions":_dimensions(original), "proposed_dimensions":_dimensions(span),
                      "removed_segments":_segment_difference(original["points"], span["points"]),
                      "added_segments":_segment_difference(span["points"], original["points"]),
                      "physical_component_ids":[item["id"] for item in physical if span["id"] in item["incident_span_ids"]],
                      "validation":_component_verification(span["id"], validation),
                      "fittings":[f for f in manufacture["fittings"] if f["span_id"] == span["id"]],
                      "manufacturing_pieces":[p for p in manufacture["pieces"] if p["span_id"] == span["id"]]})
    instructions = {"schema_version":"pipeforge.instructions/4.0", "product_version":"5.1", "status":result["status"], "proposal_id":result.get("proposal_id"),
                    "drawing":source.get("drawing", {}), "source_revision":revision, "source_digest":source_digest,
                    "units":"mm", "coordinate_system":source["coordinate_system"], "standards_profile":source.get("standards_profile"),
                    "spans":spans, "before_metrics":before, "after_metrics":after, "original_validation":original_validation,
                    "physical_components":physical,
                    "topology":{"nodes":deepcopy(candidate["nodes"]), "connections":[{"span_id":span["id"], "from_node":span["from_node"], "to_node":span["to_node"], "system_id":span["system_id"], "flow_direction":span.get("flow_direction")} for span in candidate["spans"]], "unresolved_connections":deepcopy(candidate.get("unresolved_connections", []))},
                    "environment_revision":source.get("environment_revision", 0), "environment_edits":deepcopy(source.get("environment_edits", [])),
                    "validation":validation, "unresolved_issues":unresolved, "native_apply_authorized":False, "source_input_units":source.get("input_units","mm"),
                    "validation_delta":validation_delta(original_validation, validation),
                    "network_audit":audit, "candidate_comparison":comparison, "engineering_explanation":explanation,
                    "selection":deepcopy(selection),
                    "manufacturing":manufacture, "installation_steps":_installation_steps(source, audit, manufacture, candidate_raw is not None),
                    "confidence":{"basis":"Evidence categories, not fabricated confidence percentages", "coordinate_serialization":"pending_roundtrip",
                                  "engineering_checks":audit["proposed_validation"], "commercial_and_fabrication_approval":"UNVERIFIED",
                                  "native_cad_apply":"NOT_AUTHORIZED", "installation_feasibility":"UNVERIFIED",
                                  "hydraulic_objective":after.get("hydraulic_evidence", {}).get("status", "UNVERIFIED")},
                    "implementation_state":"designer_review_required" if candidate_raw is not None else "diagnostic_or_preserve_existing"}
    metadata = {key:instructions[key] for key in ("schema_version","status","proposal_id","drawing","source_revision","source_digest","units","coordinate_system","standards_profile","implementation_state")}
    metadata["source_authority"] = "User-supplied Fincantieri challenge transcription; original authoritative PDFs and commercial approval have not been independently verified."
    metadata.update({"product_version":"5.1", "objective_config":objective_config,
                     "network_summary":{"span_count":len(audit["spans"]),"system_count":len(audit["systems"]),"changed_span_count":len(audit["changed_span_ids"]),"untouched_span_count":len(audit["untouched_span_ids"]),"moved_node_count":len(audit["moved_nodes"])},
                     "confidence":instructions["confidence"], "manufacturing_takeoff_basis":"Piece lengths include constituent bend lengths; fitting rows are requirements, not additional pipe length."})
    components = _components(candidate, manufacture)
    coordinates = _coordinates(source, candidate)
    target = Path(output_dir)
    target.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".pipeforge-export-", dir=target) as temporary:
        directory = Path(temporary)
        (directory/_EXPORT_FILES["instructions"]).write_text(_json(instructions)+"\n", encoding="utf-8")
        _write_xlsx(directory/_EXPORT_FILES["xlsx"], components, coordinates, manufacture, validation, metadata, instructions)
        with (directory/_EXPORT_FILES["csv"]).open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(_COMPONENT_COLUMNS)
            writer.writerows([[_safe_csv(value) for value in row] for row in components])
        with (directory/_EXPORT_FILES["coordinates_csv"]).open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(_COORDINATE_COLUMNS)
            writer.writerows([[_safe_csv(value) for value in row] for row in coordinates])
        _write_dxf(directory/_EXPORT_FILES["dxf"], source, candidate, manufacture)
        _write_comparison_svg(directory/_EXPORT_FILES["comparison"], source, candidate, audit)
        coordinate_check = _verify_serialized_coordinates(directory, source, candidate, coordinates, components)
        # Serialization evidence does not establish the source frame's
        # provenance. Preserve the independent WCS/units check and its unknowns.
        validation["checks"].append(coordinate_check)
        instructions["validation"] = validation
        audit["proposed_validation"] = _check_summary(validation)
        instructions["confidence"]["coordinate_serialization"] = "PASS"
        instructions["confidence"]["engineering_checks"] = audit["proposed_validation"]
        _write_xlsx(directory/_EXPORT_FILES["xlsx"], components, coordinates, manufacture, validation, metadata, instructions)
        _verify_serialized_coordinates(directory, source, candidate, coordinates, components)
        (directory/_EXPORT_FILES["instructions"]).write_text(_json(instructions)+"\n", encoding="utf-8")
        (directory/_EXPORT_FILES["instructions_markdown"]).write_text(_markdown(instructions), encoding="utf-8")
        exported_result = deepcopy(result)
        exported_result.update({"validation":validation,"manufacturing":manufacture,"export_coordinate_check":coordinate_check,
                                "physical_components":physical,"topology":instructions["topology"],
                                "validation_delta":instructions["validation_delta"],
                                "original_validation":original_validation,"network_audit":audit,"candidate_comparison":comparison,
                                "engineering_explanation":explanation,"confidence":instructions["confidence"],"installation_steps":instructions["installation_steps"],
                                "before":instructions["before_metrics"],"after":instructions["after_metrics"],
                                "changed_span_ids":[span["id"] for span in spans if span["changed"]]})
        (directory/_EXPORT_FILES["result"]).write_text(_json(exported_result)+"\n", encoding="utf-8")
        # All files are checked before exposing any of them to the caller.
        for filename in _EXPORT_FILES.values():
            (directory/filename).replace(target/filename)
    return deepcopy(_EXPORT_FILES)
