"""Revision-bound exports from synthetic global-coordinate networks."""
from copy import deepcopy
import csv
import json

import ezdxf
from openpyxl import load_workbook
import pytest

from pipe_router.connected import network_digest, metrics
from pipe_router.network_exports import export_network
from pipe_router.network_model import normalize_network


@pytest.fixture
def network():
    return normalize_network({
        "schema_version":"pipeforge.network/4.0", "id":"synthetic-export", "units":"mm", "provenance":"synthetic",
        "drawing":{"id":"synthetic-ship", "revision":"rev-4"},
        "coordinate_system":{"frame":"WCS", "up":[0,0,1], "fore_aft":[1,0,0], "athwartship":[0,1,0]},
        "bounds":{"min":[10000,20000,0], "max":[20000,30000,10000]},
        "nodes":[{"id":"equipment-a", "type":"equipment_port", "system_id":"cooling", "position":[11000,22000,3000], "cad_handle":"AA", "fixed":True},
                 {"id":"equipment-b", "type":"equipment_port", "system_id":"cooling", "position":[15000,22000,3000], "cad_handle":"BB", "fixed":True}],
        "spans":[{"id":"route-a", "from_node":"equipment-a", "to_node":"equipment-b", "system_id":"cooling", "system_type":"pressurized", "cad_handle":"CAD-123",
                  "points":[[11000,22000,3000],[15000,22000,3000]], "nominal_diameter_mm":50, "outer_diameter_mm":60,
                  "manufacturing":{"planarity_required":False,"welds":[]}}],
        "obstacles":[], "coverage":{"complete":True,"omitted_count":0,"unknown_count":0,"notes":[]},
    })


def result_for(network, candidate=None):
    return {"status":"NO IMPROVEMENT FOUND" if candidate is None else "IMPROVED", "original_network":deepcopy(network), "proposed_network":candidate,
            "source_digest":network_digest(network), "source_revision":network["drawing"]["revision"], "proposal_id":"synthetic-proposal",
            "before":metrics(network), "after":metrics(candidate or network), "unresolved_issues":[], "changed_span_ids":[]}


def test_bundle_contains_real_xlsx_exact_wcs_dxf_and_cad_identity(network,tmp_path):
    result = result_for(network)
    files = export_network(network,result,tmp_path)
    assert set(files.values()) == {"instructions.json","instructions.md","components.xlsx","components.csv","network-preview.dxf","network-result.json"}
    assert all((tmp_path/name).is_file() for name in files.values())
    instructions=json.loads((tmp_path/files["instructions"]).read_text())
    span=instructions["spans"][0]
    assert span["cad_handle"] == "CAD-123"
    assert span["original_centerline_mm"] == network["spans"][0]["points"]
    assert span["proposed_centerline_mm"] == network["spans"][0]["points"]
    assert span["fixed_connections"][0]["position"] == [11000,22000,3000]
    assert instructions["native_apply_authorized"] is False
    workbook=load_workbook(tmp_path/files["xlsx"],data_only=True)
    assert set(workbook.sheetnames) == {"Components","Coordinates","Manufacturing","Checks","Metadata"}
    coordinates=list(workbook["Coordinates"].iter_rows(min_row=2,values_only=True))
    assert coordinates[0][5:8] == (11000,22000,3000)
    assert coordinates[1][5:8] == (15000,22000,3000)
    assert coordinates[0][2] == "CAD-123"
    workbook.close()
    doc=ezdxf.readfile(tmp_path/files["dxf"])
    assert doc.units == 4
    polylines=list(doc.modelspace().query("POLYLINE"))
    assert tuple(list(polylines[0].points())[0]) == (11000,22000,3000)
    delivered=json.loads((tmp_path/files["result"]).read_text())
    assert delivered["export_coordinate_check"]["status"] == "PASS"
    assert delivered["export_coordinate_check"]["evidence"]["max_error_mm"] == 0
    assert any(item["status"] == "UNVERIFIED" for item in delivered["validation"]["checks"])


@pytest.mark.parametrize("mutation", ["revision","geometry","original"])
def test_export_rejects_stale_result_without_creating_files(network,tmp_path,mutation):
    result=result_for(network)
    if mutation == "revision":
        network["drawing"]["revision"]="rev-5"
    elif mutation == "geometry":
        network["spans"][0]["outer_diameter_mm"] = 65
    else:
        result["original_network"]["spans"][0]["outer_diameter_mm"] = 65
    with pytest.raises(ValueError, match="Stale|match"):
        export_network(network,result,tmp_path)
    assert not list(tmp_path.iterdir())


def test_export_revalidates_candidate_instead_of_trusting_result_flag(network,tmp_path):
    candidate=deepcopy(network)
    candidate["spans"][0]["points"][-1][0] += 100
    result=result_for(network,candidate)
    result["validation"]={"hard_failures":0,"checks":[]}
    with pytest.raises(ValueError,match="independent engineering validation"):
        export_network(network,result,tmp_path)
    assert not list(tmp_path.iterdir())


def test_spreadsheet_ids_are_preserved_as_strings_and_csv_is_safe(network,tmp_path):
    dangerous='=HYPERLINK("https://example.invalid")'
    network["spans"][0]["id"]=dangerous
    files=export_network(network,result_for(network),tmp_path)
    workbook=load_workbook(tmp_path/files["xlsx"],data_only=False)
    assert workbook["Coordinates"]["B2"].value == dangerous
    assert workbook["Coordinates"]["B2"].data_type == "s"
    assert workbook["Components"]["A2"].data_type == "s"
    workbook.close()
    with (tmp_path/files["csv"]).open(encoding="utf-8-sig",newline="") as stream:
        rows=list(csv.reader(stream))
    assert rows[1][0].startswith("'=")
    assert json.loads((tmp_path/files["instructions"]).read_text())["spans"][0]["id"] == dangerous


def test_export_converts_declared_meter_coordinates_to_global_mm(network,tmp_path):
    raw=deepcopy(network)
    raw.pop("input_units")
    raw["units"]="m"
    for key in ("min","max"):
        raw["bounds"][key]=[v/1000 for v in raw["bounds"][key]]
    for node in raw["nodes"]:
        node["position"]=[v/1000 for v in node["position"]]
    for span in raw["spans"]:
        span["points"]=[[v/1000 for v in point] for point in span["points"]]
    normalized=normalize_network(raw)
    export_network(raw,result_for(normalized),tmp_path)
    instructions=json.loads((tmp_path/"instructions.json").read_text())
    assert instructions["spans"][0]["proposed_centerline_mm"][0] == [11000,22000,3000]
    assert instructions["units"] == "mm"


def test_export_recomputes_metrics_and_full_length_ids(network,tmp_path):
    network["spans"][0]["id"] = "r"*256
    result=result_for(network)
    result["before"]={"length_mm":-999}
    result["after"]={"length_mm":-999}
    files=export_network(network,result,tmp_path)
    delivered=json.loads((tmp_path/files["result"]).read_text())
    assert delivered["before"]["length_mm"] == 4000
    assert delivered["after"]["length_mm"] == 4000
    assert delivered["export_coordinate_check"]["status"] == "PASS"


def test_serialized_coordinate_pass_does_not_erase_unknown_source_frame(network,tmp_path):
    network["coordinate_system"].pop("frame")
    files=export_network(network,result_for(network),tmp_path)
    delivered=json.loads((tmp_path/files["result"]).read_text())
    checks={item["id"]:item for item in delivered["validation"]["checks"]}
    assert checks["export.coordinate_consistency"]["status"] == "UNVERIFIED"
    assert checks["export.serialized_coordinate_roundtrip"]["status"] == "PASS"
