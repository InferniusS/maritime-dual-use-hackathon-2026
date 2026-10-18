"""Computed multi-system delivery, without confidential source CAD assets."""
from copy import deepcopy
import csv
import json
import math
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
import pytest

from pipe_router.network_exports import export_network
from pipe_router.network_model import network_digest, normalize_network


@pytest.fixture
def multi_system_network():
    nodes, spans = [], []
    for identifier, system, points in (
        ("cooling-a", "cooling", [[11000,22000,3000],[11000,24000,3000],[15000,24000,3000],[15000,22000,3000]]),
        ("service-b", "service", [[11000,27000,6000],[15000,27000,6000]]),
    ):
        for end, point in (("start", points[0]), ("end", points[-1])):
            nodes.append({"id": identifier+":"+end, "type":"equipment_port", "system_id":system,
                          "position":point, "fixed":True, "cad_handle":identifier+":"+end})
        spans.append({"id":identifier, "from_node":identifier+":start", "to_node":identifier+":end", "system_id":system,
                      "system_type":"pressurized", "points":points, "cad_handle":"CAD-"+identifier,
                      "nominal_diameter_mm":50, "outer_diameter_mm":60,
                      "fitting_catalog":{"radius_mm":152.4, "body_radius_mm":30, "angles_deg":[90], "status":"provisional", "source":"synthetic-test-assumption"},
                      "manufacturing":{"welds":[], "planarity_required":False}})
    return normalize_network({"schema_version":"pipeforge.network/4.0", "id":"synthetic-multi-system-export", "units":"mm", "provenance":"synthetic",
                              "drawing":{"id":"synthetic-drawing", "revision":"test-rev-5"},
                              "coordinate_system":{"frame":"WCS", "up":[0,0,1], "fore_aft":[1,0,0], "athwartship":[0,1,0]},
                              "bounds":{"min":[10000,20000,0], "max":[30000,30000,10000]},
                              "nodes":nodes, "spans":spans, "obstacles":[],
                              "coverage":{"complete":True,"omitted_count":0,"unknown_count":0,"notes":[]}})


def _result(source, candidate):
    return {"status":"IMPROVED" if candidate else "NO IMPROVEMENT FOUND", "source_revision":source["drawing"]["revision"],
            "source_digest":network_digest(source), "original_network":deepcopy(source), "proposed_network":candidate,
            "proposal_id":"synthetic-v5-export", "changed_span_ids":["invented-id"],
            "before":{"length_mm":-9999}, "after":{"length_mm":-9999}, "validation":{"hard_failures":0,"checks":[]},
            "objective_config":{"strategy":"length", "weights":{"length":1}},
            "search":{"attempts":[{"evaluation":1, "candidate_digest":"recorded-only", "metrics":{"length_mm":123456,"objective":999},
                                    "accepted":False,"hard_failures":1,"order":["cooling-a","service-b"]}]}}


def test_multi_system_bundle_recomputes_changes_tradeoffs_and_untouched_checks(multi_system_network, tmp_path):
    source = multi_system_network
    candidate = deepcopy(source)
    candidate["spans"][0]["points"] = [candidate["spans"][0]["points"][0], candidate["spans"][0]["points"][-1]]
    result = _result(source, candidate)
    export_network(source, result, tmp_path)
    instructions = json.loads((tmp_path/"instructions.json").read_text())
    audit = instructions["network_audit"]
    assert audit["changed_span_ids"] == ["cooling-a"]
    assert audit["untouched_span_ids"] == ["service-b"]
    assert instructions["before_metrics"]["length_mm"] == 12000
    assert instructions["after_metrics"]["length_mm"] == 8000
    assert instructions["before_metrics"]["normalization_scales"] == instructions["after_metrics"]["normalization_scales"]
    assert instructions["before_metrics"]["objective_strategy"] == "length"
    assert instructions["after_metrics"]["objective"] == pytest.approx(2/3)
    assert instructions["original_validation"]["checks"]
    assert audit["spans"][1]["after_validation"]["check_ids"]
    assert {row["system_id"] for row in audit["systems"]} == {"cooling","service"}
    assert audit["spans"][0]["length_delta_mm"] == -4000
    assert instructions["spans"][0]["removed_segments"][0]["start"] == [11000,22000,3000]
    assert instructions["spans"][0]["added_segments"][0]["length_mm"] == 4000
    delivered = next(row for row in instructions["candidate_comparison"] if row["candidate_id"] == "delivered")
    recorded = next(row for row in instructions["candidate_comparison"] if row["record_type"] == "optimizer_recorded_attempt")
    assert delivered["metrics"]["length_mm"] == 8000
    assert delivered["candidate_digest"] == network_digest(candidate)
    assert recorded["metrics"]["length_mm"] == 123456
    assert "Only the original and delivered" in recorded["evidence_scope"]
    assert instructions["native_apply_authorized"] is False
    assert instructions["confidence"]["coordinate_serialization"] == "PASS"
    assert instructions["confidence"]["commercial_and_fabrication_approval"] == "UNVERIFIED"
    assert any("every applicable FAIL and UNVERIFIED" in step["action"] for step in instructions["installation_steps"])
    text = (tmp_path/"instructions.md").read_text()
    assert "12000.000000 to 8000.000000" in text
    assert "service-b" in text
    assert "network-comparison.svg" in text


def test_coordinate_csv_roundtrip_and_svg_escape_untrusted_cad_ids(multi_system_network, tmp_path):
    source = multi_system_network
    source["spans"][0]["id"] = '=HYPERLINK("https://example.invalid")<script>bad</script>'
    source["spans"][0]["system_id"] = "cooling<&>"
    for node in source["nodes"][:2]:
        node["system_id"] = "cooling<&>"
    export_network(source, _result(source,None), tmp_path)
    with (tmp_path/"coordinates.csv").open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))
    assert rows[1][1].startswith("'=")
    assert [float(value) for value in rows[1][5:8]] == [11000,22000,3000]
    instructions = json.loads((tmp_path/"instructions.json").read_text())
    assert instructions["spans"][0]["id"].startswith("=")
    roundtrip = next(check for check in instructions["validation"]["checks"] if check["id"] == "export.serialized_coordinate_roundtrip")
    assert roundtrip["evidence"]["csv_coordinate_rows"] == len(rows)-1
    assert roundtrip["evidence"]["max_error_mm"] <= roundtrip["evidence"]["tolerance_mm"] == 1e-6
    svg = ET.parse(tmp_path/"network-comparison.svg").getroot()
    assert not any(element.tag.endswith("script") for element in svg.iter())
    assert not any("href" in key for element in svg.iter() for key in element.attrib)
    assert "<script>bad</script>" in "".join(svg.itertext())
    workbook = load_workbook(tmp_path/"components.xlsx", data_only=False)
    assert workbook["Changes"]["A2"].data_type == "s"
    assert workbook["Coordinates"]["B2"].data_type == "s"
    workbook.close()


def test_unique_weld_schedule_contains_exact_physical_locations(multi_system_network, tmp_path):
    source = multi_system_network
    span = source["spans"][0]
    # A supplied straight physical run needs one proposed manufacturing cut.
    span["points"] = [[11000,22000,3000],[19000,22000,3000]]
    source["nodes"][1]["position"] = [19000,22000,3000]
    export_network(source, _result(source,None), tmp_path)
    instructions = json.loads((tmp_path/"instructions.json").read_text())
    welds = [weld for weld in instructions["manufacturing"]["welds"] if weld["span_id"] == "cooling-a"]
    assert len(welds) == 1
    assert welds[0]["station_mm"] == 6096
    assert welds[0]["position"] == [17096,22000,3000]
    assert "spacing and process are separate" in welds[0]["verification_scope"]
    workbook = load_workbook(tmp_path/"components.xlsx", data_only=True)
    weld_rows = list(workbook["Welds"].iter_rows(min_row=2, values_only=True))
    assert len(weld_rows) == 1
    assert weld_rows[0][5:8] == (17096,22000,3000)
    component_rows = list(workbook["Components"].iter_rows(min_row=2, values_only=True))
    assert sum(row[3] == "weld location" for row in component_rows) == 1
    assert sum(row[6] for row in component_rows if row[3] == "manufacturing piece") == 3
    workbook.close()
    assert len(instructions["installation_steps"]) == 1
    assert instructions["installation_steps"][0]["gate"] == "diagnostic_or_unchanged"


def test_dimension_labels_are_calculated_from_three_dimensional_lengths(multi_system_network, tmp_path):
    source = multi_system_network
    span = source["spans"][1]
    span["points"][-1] = [15000,27000,9000]
    source["nodes"][3]["position"] = [15000,27000,9000]
    export_network(source, _result(source,None), tmp_path)
    svg = "".join(ET.parse(tmp_path/"network-comparison.svg").getroot().itertext())
    assert "service-b: 5000.000 mm" in svg
    instructions = json.loads((tmp_path/"instructions.json").read_text())
    assert instructions["spans"][1]["original_dimensions"]["centerline_length_mm"] == math.sqrt(4000**2+3000**2)


def test_stale_weld_or_metadata_revision_rejects_whole_bundle(multi_system_network, tmp_path):
    source = multi_system_network
    result = _result(source,None)
    source["spans"][0]["manufacturing"]["welds"] = [{"station_mm":100}]
    with pytest.raises(ValueError, match="Stale result"):
        export_network(source, result, tmp_path)
    assert not list(tmp_path.iterdir())


def test_export_rejects_physically_valid_change_outside_selected_scope(multi_system_network, tmp_path):
    from pipe_router.network_validation import validate_network
    source = multi_system_network
    candidate = deepcopy(source)
    candidate["spans"][1]["points"] = [[11000,27000,6000],[11000,28000,6000],[15000,28000,6000],[15000,27000,6000]]
    # Establish that geometry alone is acceptable; the independent permission
    # check must reject this otherwise valid change to unselected service-b.
    assert validate_network(candidate, reference=source)["hard_failures"] == 0
    result = _result(source, candidate)
    result["selection"] = {"modifiable_span_ids":["cooling-a"]}
    with pytest.raises(ValueError, match="independent engineering validation"):
        export_network(source, result, tmp_path)
    assert not list(tmp_path.iterdir())


def test_export_checks_same_selection_in_original_and_proposed_reports(multi_system_network, tmp_path):
    source = multi_system_network
    candidate = deepcopy(source)
    candidate["spans"][0]["points"] = [candidate["spans"][0]["points"][0], candidate["spans"][0]["points"][-1]]
    result = _result(source, candidate)
    result["selection"] = {"modifiable_span_ids":["cooling-a"]}
    export_network(source, result, tmp_path)
    instructions = json.loads((tmp_path/"instructions.json").read_text())
    for report in (instructions["original_validation"], instructions["validation"]):
        scope = next(check for check in report["checks"] if check["id"] == "preservation.optimization_scope")
        assert scope["status"] == "PASS"
        assert scope["evidence"]["permitted_span_ids"] == ["cooling-a"]
    assert instructions["selection"]["modifiable_span_ids"] == ["cooling-a"]
