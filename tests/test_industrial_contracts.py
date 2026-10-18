"""Lossless physical metadata, stale environments and real export readbacks."""
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
import json
import threading
import time

import ezdxf
from openpyxl import load_workbook
import pytest

from pipe_router.connected import network_digest
from pipe_router.industrial_examples import build_manifold
from pipe_router.network_components import component_geometry
from pipe_router.network_exports import export_network
from pipe_router.network_model import normalize_network, reconstruct_network
from pipe_router.network_service import NetworkRequestError, NetworkWorkspace
from test_network_api import client, network
from test_network_model import cad


def test_component_envelopes_scale_coordinates_without_scaling_mm_dimensions(network):
    network["units"] = "m"
    network["bounds"] = {"min": [-1, -1, -1], "max": [3, 3, 3]}
    network["nodes"][1]["position"] = [2, 0, 0]
    network["spans"][0]["points"] = [[0, 0, 0], [2, 0, 0]]
    network["nodes"][0].update({"type": "flange", "axis": [1, 0, 0],
        "dimensions": {"body_radius_mm": 100, "body_length_mm": 50},
        "physical_envelope": {"representation": "aabb", "min": [-.025, -.1, -.1],
                              "max": [.025, .1, .1], "anchor": [0, 0, 0],
                              "conservative": True, "source": "Synthetic supplied envelope"}})
    source = deepcopy(network)
    normalized = normalize_network(network)
    body = component_geometry(normalized)[0]
    assert body["min"] == [-25, -100, -100]
    assert body["max"] == [25, 100, 100]
    assert body["dimensions"] == {"body_radius_mm": 100, "body_length_mm": 50}
    assert body["axis"] == [1, 0, 0]
    assert normalize_network(normalized) == normalized
    assert source == network


@pytest.mark.parametrize("bad", [[], {"representation": "aabb", "min": [0, 0, 0], "max": [0, 1, 1]},
    {"representation": "capsules", "capsules": []},
    {"representation": "capsules", "capsules": [{"start": [0, 0, 0], "end": [1, 1, 1], "radius_mm": True}]}])
def test_malformed_physical_envelopes_cannot_be_normalized(network, bad):
    network["nodes"][0]["physical_envelope"] = bad
    with pytest.raises(ValueError):
        normalize_network(network)


@pytest.mark.parametrize("kind", ["component", "pipe", "fitting"])
def test_extreme_finite_dimensions_rejected_before_geometry_overflow(network, kind):
    if kind == "component":
        network["nodes"][0]["dimensions"] = {"body_radius_mm": 1e10, "body_length_mm": 1e10}
    elif kind == "pipe":
        network["spans"][0]["outer_diameter_mm"] = 1e10
    else:
        network["spans"][0]["fitting_catalog"] = {"radius_mm": 1e10}
    with pytest.raises(ValueError, match="physical geometry limit"):
        normalize_network(network)


@pytest.mark.parametrize("supports", [[{"dimensions": {"body_radius_mm": 50}}],
    [{"id": "duplicate", "type": "support"}, {"id": "duplicate", "type": "support"}]])
def test_physical_supports_require_unique_stable_identifiers(network, supports):
    network["supports"] = supports
    with pytest.raises(ValueError, match="support.*identifier"):
        normalize_network(network)


def test_unretained_cad_objects_and_failed_acis_cannot_be_reported_complete():
    source = cad()
    source["coverage"] = {"complete": True, "omitted_count": 0, "unknown_count": 0}
    source["unsupported_geometry"] = {"INSERT:external_reference": 1, "POINT": 2, "face:spline": 7, "edge:spline": 9}
    source["counts"] = {"failed_solid_instances": 1}
    result = reconstruct_network(source)
    coverage = result["coverage"]
    assert coverage["complete"] is False
    assert coverage["unsupported_entity_count"] == 3
    assert coverage["failed_solid_instances"] == 1
    assert coverage["omitted_count"] == 0  # Separate evidence, not double-counted objects.


@pytest.mark.parametrize("field,value", [("unsupported_geometry", [1]),
    ("unsupported_geometry", {"POINT": -1}), ("counts", {"failed_solid_instances": True}),
    ("counts", {"failed_solid_instances": "missing"})])
def test_malformed_cad_coverage_counters_fail_clearly(field, value):
    source = cad()
    source[field] = value
    with pytest.raises(ValueError, match="CAD"):
        reconstruct_network(source)


def test_industrial_exports_preserve_bodies_graph_and_unknown_support_coordinates(tmp_path):
    source = normalize_network(build_manifold(20))
    source["spans"][0]["pipe_id"] = "engineering-pipe-01"
    source["spans"][0]["source_layers"] = ["SYNTHETIC_COOLING"]
    source["supports"].append({"id": "unknown-location", "type": "support"})
    result = {"status": "NO IMPROVEMENT FOUND", "source_revision": source["drawing"]["revision"],
              "source_digest": network_digest(source), "original_network": deepcopy(source), "proposed_network": None}
    files = export_network(source, result, tmp_path)
    assert len(files) == 8
    instructions = json.loads((tmp_path/"instructions.json").read_text())
    assert instructions["product_version"] == "5.1"
    assert instructions["spans"][0]["pipe_id"] == "engineering-pipe-01"
    assert instructions["spans"][0]["source_layers"] == ["SYNTHETIC_COOLING"]
    assert len(instructions["topology"]["connections"]) == 20
    assert instructions["topology"]["nodes"] == source["nodes"]
    bodies = {body["id"]: body for body in instructions["physical_components"]}
    assert {body["type"] for body in bodies.values()} >= {"valve", "flange", "reducer", "support"}
    assert bodies["support:unknown-location"]["position"] is None
    assert bodies["support:unknown-location"]["status"] == "UNVERIFIED"
    for node in source["nodes"]:
        if node["id"] in bodies:
            assert bodies[node["id"]]["dimensions"] == node.get("dimensions", {})
            assert bodies[node["id"]]["position"] == node["position"]
    workbook = load_workbook(tmp_path/"components.xlsx", read_only=True)
    rows = list(workbook["Components"].iter_rows(min_row=2, values_only=True))
    unknown = next(row for row in rows if row[0] == "support:unknown-location")
    assert unknown[7:10] == (None, None, None)
    assert any(row[3] == "flange" and '"body_length_mm"' in row[5] for row in rows)
    workbook.close()
    drawing = ezdxf.readfile(tmp_path/"network-preview.dxf")
    assert len(drawing.modelspace().query('LINE[layer=="PF51_COMPONENTS"]')) > 0
    assert instructions["confidence"]["coordinate_serialization"] == "PASS"
    assert instructions["confidence"]["commercial_and_fabrication_approval"] == "UNVERIFIED"


def add_wall(network):
    return {"network": network, "operation": {"action": "add", "obstacle_id": "review-wall", "obstacle": {
        "id": "review-wall", "kind": "wall", "simulation": True,
        "min": [900, -200, -200], "max": [1100, 200, 200]}}}


def test_obstacle_http_edit_invalidates_even_original_snapshot_approval(client, network):
    calculated = client.post('/api/network/optimize', json={"network": network, "max_evaluations": 1}).json()
    edited = client.post('/api/network/obstacles/edit', json=add_wall(network))
    assert edited.status_code == 200, edited.text
    response = edited.json()
    assert response["network"]["drawing"] == network["drawing"]
    assert response["network"]["environment_revision"] == 1
    assert response["network"]["spans"] == calculated["network"]["spans"]
    assert response["affected_span_ids"] == ["ab"]
    assert response["geometry"]["spans"]
    old_approval = client.post('/api/network/review', json={"network": network, "approved": True, "review_token": calculated["review_token"]})
    assert old_approval.status_code == 409
    assert client.post('/api/network/obstacles/edit', json=add_wall(network)).status_code == 409
    assert client.post('/api/network/jobs', json={"network": network}).status_code == 409


def test_reopening_original_source_starts_new_session_without_reviving_stale_proposal(client, network):
    first = client.post('/api/network/inspect', json={"network": network}).json()['network']
    calculated = client.post('/api/network/optimize', json={"network": first, "max_evaluations": 1}).json()
    edited = client.post('/api/network/obstacles/edit', json=add_wall(first))
    assert edited.status_code == 200, edited.text
    assert edited.json()['network']['analysis_session_id'] == first['analysis_session_id']
    reopened = client.post('/api/network/inspect', json={"network": network}).json()['network']
    assert reopened['analysis_session_id'] != first['analysis_session_id']
    assert reopened['drawing'] == network['drawing']
    assert reopened['spans'] == first['spans']
    assert 'analysis_session_id' not in network
    assert client.post('/api/network/optimize', json={"network": reopened, "max_evaluations": 1}).status_code == 200
    assert client.post('/api/network/optimize', json={"network": first, "max_evaluations": 1}).status_code == 409
    assert client.post('/api/network/review', json={"network": first, "approved": True,
        "review_token": calculated['review_token']}).status_code == 409


@pytest.mark.parametrize("session", [None, [], True, 'g'*32, 'a'*31])
def test_analysis_session_identifiers_cannot_break_revision_ledger(network, session):
    network['analysis_session_id'] = session
    with pytest.raises(ValueError, match='analysis_session_id'):
        normalize_network(network)


def test_old_sync_result_cannot_republish_review_after_environment_edit(tmp_path, monkeypatch, network):
    import pipe_router.connected as connected
    real_optimizer = connected.optimize_network
    entered, release = threading.Event(), threading.Event()
    def delayed(source, **options):
        entered.set()
        assert release.wait(5)
        return real_optimizer(source, **options)
    monkeypatch.setattr(connected, "optimize_network", delayed)
    workspace = NetworkWorkspace(tmp_path)
    try:
        with ThreadPoolExecutor(max_workers=1) as pool:
            future = pool.submit(workspace.dispatch, "optimize", {"network": network, "max_evaluations": 1})
            assert entered.wait(3)
            workspace.dispatch("obstacles/edit", add_wall(network))
            release.set()
            with pytest.raises(NetworkRequestError, match="environment changed"):
                future.result(5)
        assert not workspace.reviews
        assert workspace.active == 0
    finally:
        release.set()
        workspace.close()


def test_concurrent_environment_edits_cannot_both_publish_revision_one(tmp_path, monkeypatch, network):
    import pipe_router.network_obstacles as module
    real_edit = module.edit_obstacle
    barrier = threading.Barrier(2)
    def synchronized(*args, **kwargs):
        result = real_edit(*args, **kwargs)
        barrier.wait(5)
        return result
    monkeypatch.setattr(module, "edit_obstacle", synchronized)
    workspace = NetworkWorkspace(tmp_path)
    try:
        with ThreadPoolExecutor(max_workers=2) as pool:
            futures = [pool.submit(workspace.dispatch, "obstacles/edit", add_wall(network)) for _ in range(2)]
            successes, conflicts = [], []
            for future in futures:
                try:
                    successes.append(future.result(6))
                except NetworkRequestError as error:
                    conflicts.append(error.status)
        assert len(successes) == 1
        assert conflicts == [409]
    finally:
        workspace.close()
